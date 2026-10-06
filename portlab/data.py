"""Offline data contracts, exploratory annual forecasts, and auditable reports.

Annual reported units never become vessel schedules.  Uploaded operational files
are unverified unless their rows explicitly carry observation provenance.
"""
from __future__ import annotations

from dataclasses import asdict
from datetime import date, datetime, timezone
from pathlib import Path
import csv
import html
import json
import math
import statistics
from typing import Any

from .contracts import DatasetBundle, SimulationResult, VesselCall


ALIASES = {
    'vessel_id': ('vessel_id', 'VesselID', 'call_id'),
    'arrival_hour': ('arrival_hour', 'ETA_RealHours', 'eta_real_hour', 'arrival_real_hour'),
    'arrival_utc': ('arrival_utc', 'eta_utc'),
    'import_boxes': ('import_boxes', 'ImportContainers', 'import_containers'),
    'export_boxes': ('export_boxes', 'ExportContainers', 'export_containers'),
    'length_m': ('length_m', 'vessel_length_m', 'LengthMeters', 'LOA_m'),
    'draft_m': ('draft_m', 'vessel_draft_m', 'DraftMeters'),
    'box40_share': ('box40_share', 'share_40ft', 'fraction40'),
    'port': ('port', 'terminal'),
    'fiscal_year': ('fiscal_year', 'year', 'FY'),
    'metric': ('metric', 'measure'),
    'value': ('value', 'annual_volume'),
    'unit': ('unit', 'units', 'cargo_unit'),
    'record_status': ('record_status', 'status'),
    'evidence_type': ('evidence_type', 'provenance'),
    'evidence_reference': ('evidence_reference', 'source_id', 'source_reference'),
    'note': ('note', 'notes'),
    'grain': ('grain', 'frequency', 'period_type'),
    'period_start': ('period_start', 'start_date'),
    'period_end': ('period_end', 'end_date'),
    'berth_start_utc': ('berth_start_utc',),
    'departure_utc': ('departure_utc',),
    'observed_berth_start_hour': ('observed_berth_start_hour', 'observed_berth_start_real_hour'),
    'observed_departure_hour': ('observed_departure_hour', 'observed_departure_real_hour'),
    'expected_departure_hour': ('expected_departure_hour', 'ExpectedDepartureRealHour'),
    'planned_stay_hours': ('planned_stay_hours', 'PlannedPortStayRealHours'),
}


def _key(value: Any) -> str:
    return ''.join(char.lower() for char in str(value or '') if char.isalnum())


_COLUMN_MAP = {_key(alias): name for name, aliases in ALIASES.items() for alias in aliases}


def _port(value: Any) -> str:
    text = str(value or 'Guam').strip()
    lookup = {'guam': 'Guam', 'portofguam': 'Guam', 'conley': 'Conley',
              'conleyterminal': 'Conley', 'bostonconley': 'Conley'}
    return lookup.get(_key(text), text)


def _number(value: Any, field: str, context: str, integer: bool = False,
            minimum: float = 0) -> float | int:
    if isinstance(value, bool) or value is None or str(value).strip() == '':
        raise ValueError(f'{context}: {field} must be a number; blank is not zero')
    try:
        parsed = float(str(value).replace(',', '').strip())
    except (ValueError, TypeError):
        raise ValueError(f'{context}: invalid {field} {value!r}') from None
    if not math.isfinite(parsed) or parsed < minimum:
        raise ValueError(f'{context}: {field} must be finite and >= {minimum:g}')
    if integer and not parsed.is_integer():
        raise ValueError(f'{context}: {field} must be a whole box count or integer')
    return int(parsed) if integer else parsed


def _timestamp(value: Any, context: str, utc: bool = False) -> datetime:
    if isinstance(value, datetime):
        stamp = value
    elif isinstance(value, date):
        stamp = datetime.combine(value, datetime.min.time())
    else:
        try:
            stamp = datetime.fromisoformat(str(value).strip().replace('Z', '+00:00'))
        except (ValueError, TypeError):
            raise ValueError(f'{context}: use ISO dates/times, e.g. 2026-10-04T12:00:00Z') from None
    # A column explicitly named UTC supplies the timezone for Excel datetimes.
    if utc:
        if stamp.tzinfo is None:
            stamp = stamp.replace(tzinfo=timezone.utc)
        return stamp.astimezone(timezone.utc)
    return stamp.replace(tzinfo=None)


def _tables(path: Path):
    """Yield (sheet name, normalized records). Never execute macros or formulas."""
    if path.suffix.lower() in ('.csv', '.tsv'):
        with path.open('r', encoding='utf-8-sig', newline='') as handle:
            rows = csv.reader(handle, delimiter='\t' if path.suffix.lower() == '.tsv' else ',')
            yield path.stem, _records(rows, path.name)
    elif path.suffix.lower() == '.xlsx':
        try:
            import openpyxl
        except ImportError:
            raise ValueError('XLSX import needs the bundled openpyxl dependency') from None
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in book.worksheets:
                yield sheet.title, _records(sheet.iter_rows(values_only=True), f'{path.name}/{sheet.title}')
        finally:
            book.close()
    else:
        raise ValueError(f'{path.name}: supported files are CSV, TSV and XLSX (save old XLS as XLSX)')


def _records(rows, context: str) -> list[dict]:
    iterator = iter(rows)
    header = None
    start = 0
    # Permit short titles above the header, but never interpret arbitrary cells as data.
    for number, row in enumerate(iterator, 1):
        mapped = [_COLUMN_MAP.get(_key(value)) for value in row]
        if ('vessel_id' in mapped and ('arrival_hour' in mapped or 'arrival_utc' in mapped)) or ('fiscal_year' in mapped and 'value' in mapped):
            header, start = mapped, number
            break
        if number >= 40:
            return []
    if header is None:
        return []
    keys = [name for name in header if name is not None]
    if len(keys) != len(set(keys)):
        raise ValueError(f'{context}: duplicate or conflicting column aliases')
    result = []
    for number, row in enumerate(iterator, start + 1):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        record = {key: row[index] for index, key in enumerate(header)
                  if key is not None and index < len(row)}
        record['_context'] = f'{context} row {number}'
        result.append(record)
        if len(result) > 250000:
            raise ValueError(f'{context}: more than 250,000 records; split this upload')
    return result


def _provenance(row: dict, annual: bool = False) -> tuple[str, str]:
    context = row['_context']
    stated = str(row.get('evidence_type') or row.get('record_status') or '').strip().lower()
    reference = str(row.get('evidence_reference') or '').strip()
    if stated in ('synthetic', 'illustrative', 'example', 'simulated', 'demo'):
        return 'synthetic', reference
    if stated in ('forecast', 'forecast_estimate', 'estimated', 'estimate', 'projection', 'planned'):
        return 'estimate', reference
    if stated in ('observed', 'actual', 'official_port_report', 'official_acfr', 'reported_unaudited'):
        if not reference:
            raise ValueError(f'{context}: observed records need evidence_reference/source_id')
        return 'observed', reference
    if stated in ('', 'unverified', 'uploaded', 'unknown'):
        return 'unverified', reference
    raise ValueError(f'{context}: unsupported evidence_type/record_status {stated!r}')


def _unit(value: Any, context: str) -> str:
    units = {'box': 'boxes', 'boxes': 'boxes', 'container': 'boxes', 'containers': 'boxes',
             'teu': 'TEU', 'teus': 'TEU'}
    match = units.get(_key(value))
    if not match:
        raise ValueError(f'{context}: explicit unit must be boxes or TEU; units cannot be inferred')
    return match


def load_datasets(paths: list[str | Path], default_port: str = 'Guam') -> DatasetBundle:
    bundle = DatasetBundle()
    raw_calls = []
    annual_keys = set()
    call_keys = set()
    for input_path in paths:
        path = Path(input_path)
        if not path.is_file():
            raise ValueError(f'Dataset not found: {path}')
        bundle.sources.append(str(path.resolve()))
        recognized = False
        for sheet, rows in _tables(path):
            if not rows:
                continue
            if 'vessel_id' in rows[0]:
                recognized = True
                for row in rows:
                    context = row['_context']
                    stated_port = str(row.get('port') or '').strip()
                    port = _port(stated_port or default_port)
                    if not stated_port:
                        bundle.evidence_notes.append(f'{path.name}/{sheet}: unlabeled AnyLogic schedule assigned to selected port {port}; confirm the port.')
                    vessel_id = str(row.get('vessel_id') or '').strip()
                    if not vessel_id:
                        raise ValueError(f'{context}: vessel_id is required')
                    identity = (port.casefold(), vessel_id.casefold())
                    if identity in call_keys:
                        raise ValueError(f'{context}: duplicate vessel/call ID {port}/{vessel_id}; repeated visits need distinct call IDs')
                    call_keys.add(identity)
                    for name in ('import_boxes', 'export_boxes'):
                        if name not in row:
                            raise ValueError(f'{context}: {name} is required. TEU cannot be converted automatically into boxes.')
                    if row.get('unit') is not None and _unit(row['unit'], context) != 'boxes':
                        raise ValueError(f'{context}: vessel workload must be boxes, not TEU')
                    import_boxes = _number(row.get('import_boxes'), 'import_boxes', context, True)
                    export_boxes = _number(row.get('export_boxes'), 'export_boxes', context, True)
                    has_hour = row.get('arrival_hour') not in (None, '')
                    has_stamp = row.get('arrival_utc') not in (None, '')
                    if has_hour == has_stamp:
                        raise ValueError(f'{context}: provide exactly one of arrival_hour/ETA_RealHours or arrival_utc')
                    arrival = (_number(row['arrival_hour'], 'arrival_hour', context) if has_hour
                               else _timestamp(row['arrival_utc'], context, utc=True))
                    evidence, reference = _provenance(row)
                    if evidence == 'estimate':
                        evidence = 'unverified'
                    measured = any(row.get(name) not in (None, '') for name in
                                   ('berth_start_utc', 'departure_utc', 'observed_berth_start_hour', 'observed_departure_hour'))
                    if measured and evidence != 'observed':
                        raise ValueError(f'{context}: measured service times require evidence_type=observed and an evidence_reference; remove illustrative times or supply observation provenance')
                    length = _number(row.get('length_m') if row.get('length_m') not in (None, '') else 180,
                                     'length_m', context, minimum=0.01)
                    draft = _number(row.get('draft_m') if row.get('draft_m') not in (None, '') else 8,
                                    'draft_m', context, minimum=0.01)
                    share = _number(row.get('box40_share') if row.get('box40_share') not in (None, '') else .5,
                                    'box40_share', context)
                    if share > 1:
                        raise ValueError(f'{context}: box40_share must be between 0 and 1')
                    for time_key in ('berth_start_utc', 'departure_utc'):
                        if row.get(time_key) not in (None, ''):
                            if has_hour:
                                raise ValueError(f'{context}: observed UTC service timestamps require arrival_utc')
                            row[time_key] = _timestamp(row[time_key], context, utc=True)
                            if row[time_key] < arrival:
                                raise ValueError(f'{context}: {time_key} precedes arrival_utc')
                    for time_key in ('observed_berth_start_hour', 'observed_departure_hour'):
                        if row.get(time_key) not in (None, ''):
                            if not has_hour:
                                raise ValueError(f'{context}: observed elapsed-hour service times require arrival_hour; use *_utc fields with arrival_utc')
                            row[time_key] = _number(row[time_key], time_key, context)
                            if row[time_key] < arrival:
                                raise ValueError(f'{context}: {time_key} precedes arrival_hour')
                    if (isinstance(row.get('berth_start_utc'), datetime) and isinstance(row.get('departure_utc'), datetime)
                            and row['departure_utc'] < row['berth_start_utc']):
                        raise ValueError(f'{context}: departure precedes berth start')
                    if (row.get('observed_berth_start_hour') not in (None, '') and row.get('observed_departure_hour') not in (None, '')
                            and row['observed_departure_hour'] < row['observed_berth_start_hour']):
                        raise ValueError(f'{context}: observed departure precedes berth start')
                    for time_key in ('expected_departure_hour', 'planned_stay_hours'):
                        if row.get(time_key) not in (None, ''):
                            checked = _number(row[time_key], time_key, context)
                            if time_key == 'expected_departure_hour' and has_hour and checked < arrival:
                                raise ValueError(f'{context}: expected departure precedes arrival')
                    if reference:
                        bundle.evidence_notes.append(f'{port}/{vessel_id}: {evidence}; reference {reference}.')
                    if row.get('note'):
                        bundle.evidence_notes.append(f'{port}/{vessel_id}: {row["note"]}')
                    call = VesselCall(vessel_id, 0, import_boxes, export_boxes, length, draft, share, port, evidence)
                    if has_hour:
                        call.observed_berth_start_hour = row.get('observed_berth_start_hour') if row.get('observed_berth_start_hour') not in (None, '') else None
                        call.observed_departure_hour = row.get('observed_departure_hour') if row.get('observed_departure_hour') not in (None, '') else None
                    else:
                        # Retain UTC until all selected files establish one common origin.
                        call.observed_berth_start_hour = row.get('berth_start_utc') if row.get('berth_start_utc') not in (None, '') else None
                        call.observed_departure_hour = row.get('departure_utc') if row.get('departure_utc') not in (None, '') else None
                    raw_calls.append((call, arrival))
            elif 'fiscal_year' in rows[0] and 'value' in rows[0]:
                recognized = True
                for row in rows:
                    context = row['_context']
                    grain = _key(row.get('grain') or 'annual')
                    if grain not in ('annual', 'annually', 'year', 'yearly', 'fiscalyear', 'fy'):
                        bundle.evidence_notes.append(f'{context}: {grain} row excluded from annual forecasts; no aggregation performed.')
                        continue
                    if not str(row.get('port') or '').strip():
                        raise ValueError(f'{context}: annual history needs an explicit port')
                    port = _port(row['port'])
                    year = _number(row.get('fiscal_year'), 'fiscal_year', context, True, 1900)
                    if year > 2200:
                        raise ValueError(f'{context}: fiscal_year exceeds 2200')
                    metric = _unit(row.get('metric') or row.get('unit'), context)
                    unit = _unit(row.get('unit') or row.get('metric'), context)
                    if metric != unit:
                        raise ValueError(f'{context}: metric and unit disagree')
                    if row.get('period_start') not in (None, '') or row.get('period_end') not in (None, ''):
                        if row.get('period_start') in (None, '') or row.get('period_end') in (None, ''):
                            raise ValueError(f'{context}: supply both period_start and period_end')
                        start = _timestamp(row['period_start'], context)
                        end = _timestamp(row['period_end'], context)
                        duration = (end - start).total_seconds() / 86400
                        if duration < 360 or duration > 367:
                            bundle.evidence_notes.append(f'{context}: nonannual period ({duration:g} days) excluded from annual forecasts.')
                            continue
                        if end.year != year:
                            raise ValueError(f'{context}: fiscal_year must match period_end year')
                    value = _number(row.get('value'), 'value', context, True)
                    evidence, reference = _provenance(row, annual=True)
                    if evidence in ('synthetic', 'estimate'):
                        bundle.evidence_notes.append(f'{context}: {evidence} annual value excluded from historical forecasting.')
                        continue
                    identity = (port.casefold(), metric, year)
                    if identity in annual_keys:
                        raise ValueError(f'{context}: duplicate annual value {port}/{metric}/FY{year}; reconcile revisions before upload')
                    annual_keys.add(identity)
                    bundle.annual_series.setdefault((port, metric), []).append((year, value))
                    bundle.evidence_notes.append(f'{port} FY{year} {value:,} {metric}: {evidence}'
                                                 + (f'; source {reference}' if reference else '; uploaded without observation provenance')
                                                 + (f'. {row["note"]}' if row.get('note') else '.'))
        if not recognized:
            raise ValueError(f'{path.name}: no recognized vessel-call or annual-history table. See templates.')
    for port in {call.port for call, _ in raw_calls}:
        subset = [(call, arrival) for call, arrival in raw_calls if call.port == port]
        times = [arrival for _, arrival in subset]
        absolute = [isinstance(arrival, datetime) for arrival in times]
        if any(absolute) and not all(absolute):
            raise ValueError(f'{port}: cannot combine relative hours and UTC dates without an explicit shared origin; upload one time basis')
        if all(absolute):
            origin = min(times)
            bundle.evidence_notes.append(f'{port}: hour zero is {origin.isoformat()}; naive values in *_utc columns are interpreted as UTC.')
            for call, arrival in subset:
                call.arrival_hour = (arrival - origin).total_seconds() / 3600
                for name in ('observed_berth_start_hour', 'observed_departure_hour'):
                    stamp = getattr(call, name)
                    if isinstance(stamp, datetime):
                        setattr(call, name, (stamp - origin).total_seconds() / 3600)
        else:
            for call, arrival in subset:
                call.arrival_hour = arrival
        bundle.calls.extend(call for call, _ in subset)
    bundle.calls.sort(key=lambda call: (call.port, call.arrival_hour, call.vessel_id))
    for values in bundle.annual_series.values():
        values.sort()
    if bundle.calls:
        counts = {kind: sum(call.evidence_type == kind for call in bundle.calls)
                  for kind in ('observed', 'unverified', 'synthetic')}
        bundle.evidence_notes.append('Vessel-call provenance: ' + ', '.join(f'{count} {kind}' for kind, count in counts.items()) + '.')
        bundle.evidence_notes.append('Missing dimensions use 180 m length and 8 m draft as assumptions. Supplied dimensions are checked against settings, not bathymetry.')
        bundle.evidence_notes.append('40-foot share defaults to 0.5 when absent. This is a sensitivity assumption, not a measured Guam or Conley size mix. 3D asset size does not establish cargo composition.')
        bundle.evidence_notes.append('Explicit observed service timestamps are retained for measured-versus-simulated timing checks and are not used to fit the DES. Such checks alone do not validate the complete port model. Simulated wait and service remain conditional outputs.')
    if bundle.annual_series:
        bundle.evidence_notes.append('Annual boxes and TEU are modeled separately. Annual history does not generate vessel arrivals, within-year seasonality or cargo per call.')
    bundle.evidence_notes = list(dict.fromkeys(bundle.evidence_notes))
    return bundle


def load_demo(data_dir: Path) -> DatasetBundle:
    bundle = load_datasets([data_dir / 'illustrative_calls.csv', data_dir / 'Port_Guam_Conley_History.xlsx'])
    bundle.evidence_notes.extend([
        'Illustrative operational workload generated with Python random.Random(42); it is independent of observed annual totals and is designed to exercise congestion and recommendation screens.',
        'Historical transcription is frozen as supplied on 2026-10-04, through FY2025. No FY2026 actual is included; the next-period forecast is FY2026, even though its fiscal period has ended.',
        'Conley FY2018 TEU 281,978 and FY2020 TEU 283,061 follow FY2025 ACFR S-15. Older reports differ (283,720 and 282,629 respectively); cause unresolved.',
        'Guam fiscal year ends September 30; Conley fiscal year ends June 30. Guam FY2025 83,574 boxes is a board-reported unaudited figure.',
        'Equipment snapshots and project geometry are dated evidence, not staffed resource availability or operational calibration. All demo service times, tractor cycles, dwell and capacities are editable assumptions.',
        'Incremental labor/energy/capital costs and avoidable vessel delay values are absent. No financial ROI or procurement ranking is supported.',
    ])
    register_path = data_dir / 'source_register.json'
    if register_path.is_file():
        for source in json.loads(register_path.read_text(encoding='utf-8')):
            bundle.evidence_notes.append(f'Source {source["source_id"]}: {source["organization"]}. {source["document"]}. '
                                         f'{source.get("location / unit definition", "")}. {source.get("url") or "User archive"}. '
                                         f'{source.get("source note", "")}')
    return bundle


def dataset_from_project(value: Any) -> DatasetBundle:
    """Validate schema-1 saved JSON as data without reading paths or executing it.

    Settings are validated separately by TerminalConfig. No conversions from
    TEU to boxes, missing observations to zero, or fractional years to integers
    are performed when restoring a project.
    """
    if type(value) is not dict:
        raise ValueError('Project must be a JSON object containing schema_version, calls and annual_series')
    if type(value.get('schema_version')) is not int or value['schema_version'] != 1:
        raise ValueError('Unsupported project schema_version; this application opens schema 1 projects')
    known_fields = {'schema_version', 'config', 'calls', 'annual_series', 'evidence_notes', 'sources'}
    unknown_fields = set(value) - known_fields
    if unknown_fields:
        raise ValueError('Unsupported project fields: ' + ', '.join(sorted(map(str, unknown_fields))))

    def required_list(key: str, default=None):
        collection = value.get(key, default)
        if type(collection) is not list:
            raise ValueError(f'Project {key} must be a JSON list')
        return collection

    def text(item: Any, name: str, context: str, blank: bool = False):
        if not isinstance(item, str) or (not blank and not item.strip()):
            raise ValueError(f'{context}: {name} must be ' + ('text' if blank else 'nonblank text'))
        return item if blank else item.strip()

    def finite(item: Any, name: str, context: str, minimum: float = 0):
        valid = type(item) in (int, float)
        if valid:
            try:
                valid = math.isfinite(item) and item >= minimum
            except (OverflowError, TypeError, ValueError):
                valid = False
        if not valid:
            raise ValueError(f'{context}: {name} must be a finite JSON number >= {minimum:g}; strings, booleans and NaN are invalid')
        return item

    def whole(item: Any, name: str, context: str, minimum: int = 0):
        if type(item) is not int or item < minimum:
            raise ValueError(f'{context}: {name} must be a whole JSON integer >= {minimum}; booleans and fractional counts are invalid')
        finite(item, name, context, minimum)
        return item

    bundle = DatasetBundle()
    calls = required_list('calls')
    if len(calls) > 250000:
        raise ValueError('Project contains more than 250,000 calls; split the project before opening')
    optional_call_fields = {'observed_berth_start_hour', 'observed_departure_hour'}
    call_fields = set(VesselCall.__dataclass_fields__)
    seen_calls = set()
    for index, record in enumerate(calls, 1):
        context = f'Project call {index}'
        if type(record) is not dict:
            raise ValueError(f'{context}: each call must be a JSON object')
        missing = call_fields - optional_call_fields - set(record)
        unknown = set(record) - call_fields
        if missing:
            raise ValueError(f'{context}: missing fields ' + ', '.join(sorted(missing)))
        if unknown:
            raise ValueError(f'{context}: unsupported fields ' + ', '.join(sorted(map(str, unknown))))
        identifier = text(record['vessel_id'], 'vessel_id', context)
        port = _port(text(record['port'], 'port', context))
        identity = (port.casefold(), identifier.casefold())
        if identity in seen_calls:
            raise ValueError(f'{context}: duplicate vessel/call ID {port}/{identifier}; repeated visits need distinct call IDs')
        seen_calls.add(identity)
        arrival = finite(record['arrival_hour'], 'arrival_hour', context)
        imports = whole(record['import_boxes'], 'import_boxes', context)
        exports = whole(record['export_boxes'], 'export_boxes', context)
        length = finite(record['length_m'], 'length_m', context, .000001)
        draft = finite(record['draft_m'], 'draft_m', context, .000001)
        share = finite(record['box40_share'], 'box40_share', context)
        if share > 1:
            raise ValueError(f'{context}: box40_share must be a fraction between 0 and 1')
        evidence = text(record['evidence_type'], 'evidence_type', context).casefold()
        if evidence not in ('observed', 'unverified', 'synthetic'):
            raise ValueError(f'{context}: evidence_type must be observed, unverified or synthetic; provenance labels are data declarations')
        observed = {}
        for name in optional_call_fields:
            measurement = record.get(name)
            if measurement is not None:
                measurement = finite(measurement, name, context)
                if evidence != 'observed':
                    raise ValueError(f'{context}: {name} requires observed provenance')
                if measurement < arrival:
                    raise ValueError(f'{context}: {name} precedes arrival_hour')
            observed[name] = measurement
        if (all(observed[name] is not None for name in optional_call_fields)
                and observed['observed_departure_hour'] < observed['observed_berth_start_hour']):
            raise ValueError(f'{context}: observed departure precedes berth start')
        bundle.calls.append(VesselCall(identifier, arrival, imports, exports, length, draft, share, port, evidence, **observed))
    series = required_list('annual_series')
    seen_series = set()
    for index, record in enumerate(series, 1):
        context = f'Project annual series {index}'
        if type(record) is not dict:
            raise ValueError(f'{context}: each series must be a JSON object')
        missing = {'port', 'metric', 'history'} - set(record)
        unknown = set(record) - {'port', 'metric', 'history', 'unit'}
        if missing:
            raise ValueError(f'{context}: missing fields ' + ', '.join(sorted(missing)))
        if unknown:
            raise ValueError(f'{context}: unsupported fields ' + ', '.join(sorted(map(str, unknown))))
        port = _port(text(record['port'], 'port', context))
        metric = _unit(text(record['metric'], 'metric', context), context)
        if 'unit' in record and _unit(text(record['unit'], 'unit', context), context) != metric:
            raise ValueError(f'{context}: metric and unit disagree; boxes and TEU remain separate')
        identity = (port.casefold(), metric)
        if identity in seen_series:
            raise ValueError(f'{context}: duplicate annual series {port}/{metric}; reconcile revisions before opening')
        seen_series.add(identity)
        history = record['history']
        if type(history) is not list:
            raise ValueError(f'{context}: history must be a list of [fiscal_year, value] pairs')
        years = set()
        clean = []
        for row_index, pair in enumerate(history, 1):
            row_context = f'{context} history row {row_index}'
            if type(pair) not in (list, tuple) or len(pair) != 2:
                raise ValueError(f'{row_context}: expected [fiscal_year, value]')
            year = whole(pair[0], 'fiscal_year', row_context, 1900)
            if year > 2200:
                raise ValueError(f'{row_context}: fiscal_year exceeds 2200')
            observed_value = finite(pair[1], 'annual value', row_context)
            if year in years:
                raise ValueError(f'{row_context}: duplicate fiscal year {year} for {port}/{metric}')
            years.add(year)
            clean.append((year, observed_value))
        bundle.annual_series[(port, metric)] = sorted(clean)
    for name in ('evidence_notes', 'sources'):
        entries = required_list(name, [])
        for index, entry in enumerate(entries, 1):
            text(entry, name, f'Project {name} entry {index}', blank=True)
        setattr(bundle, name, list(entries))
    bundle.calls.sort(key=lambda call: (call.port.casefold(), call.arrival_hour, call.vessel_id.casefold()))
    return bundle


_METHODS = ('last_year', 'mean3', 'damped_trend')
_MIN_INNER_HOLDOUTS = 3
_PROMOTION_MARGIN = .05
_RECENT_INNER_WINS = 2
_METHOD_NOTES = {
    'last_year': 'Last observed annual value (naive benchmark).',
    'mean3': 'Mean of the last three annual observations.',
    'damped_trend': 'Custom conservative trend: last value plus half the ordinary-least-squares slope over at most five annual observations; clipped at zero. This is not fitted Holt damped exponential smoothing.',
}


def _predict(values: list[float], method: str) -> float:
    if method == 'last_year':
        return float(values[-1])
    if method == 'mean3':
        return statistics.mean(values[-3:])
    if method == 'damped_trend':
        recent = values[-5:]
        xbar = (len(recent) - 1) / 2
        denominator = sum((index - xbar) ** 2 for index in range(len(recent)))
        slope = sum((index - xbar) * (value - statistics.mean(recent)) for index, value in enumerate(recent)) / denominator if denominator else 0
        return max(0.0, values[-1] + .5 * slope)
    raise ValueError(f'Unknown forecast method: {method}')


def _select(values: list[float]) -> str:
    return _selection_details(values)['selected_method']


def _selection_details(values: list[float]) -> dict:
    """Promote a challenger only with enough earlier, recent benchmark wins.

    Three folds, a 5% MAE margin, and two most-recent wins are transparent
    prototype safeguards, not statistically established selection thresholds.
    Every comparison uses targets strictly before the outer/future target.
    """
    folds = list(range(3, len(values)))
    errors = {method: [abs(_predict(values[:cut], method) - values[cut]) for cut in folds]
              for method in _METHODS}
    losses = {method: statistics.fmean(error) if error else None for method, error in errors.items()}
    recent_wins = {method: (len(folds) >= _RECENT_INNER_WINS and
                           all(candidate < baseline for candidate, baseline in
                               zip(errors[method][-_RECENT_INNER_WINS:], errors['last_year'][-_RECENT_INNER_WINS:])))
                   for method in _METHODS if method != 'last_year'}
    details = {'selected_method': 'last_year', 'inner_holdouts': len(folds),
               'candidate_mae': losses, 'recent_inner_wins': recent_wins,
               'minimum_inner_holdouts': _MIN_INNER_HOLDOUTS,
               'minimum_relative_mae_improvement': _PROMOTION_MARGIN,
               'required_latest_inner_wins': _RECENT_INNER_WINS,
               'threshold_status': 'heuristic_prototype_guards'}
    if len(folds) < _MIN_INNER_HOLDOUTS:
        details['reason'] = f'Keep last-year baseline: only {len(folds)} inner holdout(s); at least {_MIN_INNER_HOLDOUTS} required to promote a challenger.'
        return details
    best = min(_METHODS, key=lambda method: (losses[method], _METHODS.index(method)))
    baseline_mae = losses['last_year']
    if best == 'last_year' or baseline_mae == 0:
        details['reason'] = 'Keep last-year baseline: no challenger has lower earlier inner-holdout MAE.'
    elif losses[best] > baseline_mae * (1 - _PROMOTION_MARGIN):
        details['reason'] = 'Keep last-year baseline: challenger improvement is smaller than the prototype 5% MAE margin.'
    elif not recent_wins[best]:
        details['reason'] = 'Keep last-year baseline: best average challenger did not beat it in both latest inner holdouts.'
    else:
        details['selected_method'] = best
        details['reason'] = f'Promote {best}: at least three earlier inner holdouts, at least 5% lower MAE, and wins in both latest inner holdouts. Heuristic screen only.'
    return details


def _accuracy(tests: list[dict]) -> dict:
    if not tests:
        return {'holdouts': 0, 'mae': None, 'rmse': None, 'mean_error': None, 'wape_percent': None}
    errors = [test['forecast'] - test['actual'] for test in tests]
    denominator = sum(abs(test['actual']) for test in tests)
    return {'holdouts': len(tests), 'mae': statistics.fmean(abs(error) for error in errors),
            'rmse': math.sqrt(statistics.fmean(error * error for error in errors)),
            'mean_error': statistics.fmean(errors),
            'wape_percent': 100 * sum(abs(error) for error in errors) / denominator if denominator else None}


def forecast_series(annual_series: dict) -> list[dict]:
    """Nested one-year rolling forecasts, fixed-method benchmarks and errors.

    The outer score evaluates the guarded selection policy, not just the
    production method selected using the latest history. Method-comparison
    rows use identical outer years and are descriptive retrospective checks.
    """
    output = []
    for (port, metric), series in sorted(annual_series.items()):
        if not isinstance(port, str) or not port.strip():
            raise ValueError('Annual forecast port must be nonblank text')
        metric = _unit(metric, f'{port} annual forecast')
        for pair in series:
            if (not isinstance(pair, (list, tuple)) or len(pair) != 2 or
                    type(pair[0]) is not int or not 1900 <= pair[0] <= 2200):
                raise ValueError(f'{port}/{metric}: each annual observation needs a whole fiscal year and numeric value')
            if type(pair[1]) not in (int, float) or not math.isfinite(pair[1]) or pair[1] < 0:
                raise ValueError(f'{port}/{metric}: annual values must be finite and nonnegative JSON numbers')
        pairs = sorted(series)
        if not pairs:
            continue
        years = [year for year, _ in pairs]
        values = [float(value) for _, value in pairs]
        if len(set(years)) != len(years):
            raise ValueError(f'{port}/{metric}: duplicate fiscal years')
        common = {'port': port.strip(), 'metric': metric, 'unit': metric, 'years': len(pairs),
                  'history': [{'year': year, 'value': value} for year, value in pairs],
                  'history_end_year': years[-1], 'forecast_year': years[-1] + 1,
                  'forecast_horizon_years': 1, 'policy_method': 'guarded_baseline_selection',
                  'evaluation_note': 'Retrospective nested rolling-origin replay: every target was hidden from its fit and method selection. This development rerun is not independent prospective validation.',
                  'error_sign_note': 'Error = forecast minus actual; positive mean error indicates overprediction.',
                  'metric_note': 'MAE/RMSE are in the series unit. WAPE is total absolute error divided by total heldout actual volume; it is unavailable if every heldout actual is zero and is not an accuracy probability.',
                  'baseline_comparison_note': 'Fixed methods and guarded policy are scored on identical outer fiscal years. Retrospective benchmark rankings do not select the production method.',
                  'method_notes': _METHOD_NOTES.copy()}
        contiguous = all(right == left + 1 for left, right in zip(years, years[1:]))
        if not contiguous:
            output.append({**common, 'forecast': None, 'method': None,
                           'status': 'insufficient_data', **_accuracy([]),
                           'tested_first_year': None, 'tested_last_year': None,
                           'baseline_wape_percent': None, 'baseline_mae': None, 'baseline_skill_percent': None,
                           'selection_details': {}, 'method_comparison': [],
                           'envelope_evaluated_holdouts': 0, 'envelope_hits': 0, 'envelope_hit_rate_percent': None,
                           'uncertainty_low': None, 'uncertainty_high': None,
                           'uncertainty_note': 'No forecast: annual gaps must be resolved before annual rolling holdouts.',
                           'reasons': ['Nonconsecutive fiscal years. Missing observations are not zero.'], 'backtests': []})
            continue
        tests = []
        fixed_tests = {method: [] for method in _METHODS}
        # Outer targets never participate in choosing their method.
        for cut in range(4, len(values)):
            details = _selection_details(values[:cut])
            method = details['selected_method']
            prediction = _predict(values[:cut], method)
            for fixed in _METHODS:
                fixed_tests[fixed].append({'year': years[cut], 'actual': values[cut],
                                           'forecast': _predict(values[:cut], fixed)})
            prior_envelope = max((abs(test['error']) for test in tests), default=None)
            lower = max(0, prediction - prior_envelope) if prior_envelope is not None else None
            upper = prediction + prior_envelope if prior_envelope is not None else None
            baseline = _predict(values[:cut], 'last_year')
            tests.append({'year': years[cut], 'actual': values[cut], 'forecast': prediction,
                          'method': method, 'error': prediction - values[cut],
                          'absolute_error': abs(prediction - values[cut]),
                          'absolute_percentage_error': 100 * abs(prediction - values[cut]) / values[cut] if values[cut] else None,
                          'train_last_year': years[cut - 1], 'train_years': cut,
                          'naive_forecast': baseline, 'naive_error': baseline - values[cut],
                          'inner_holdouts': details['inner_holdouts'], 'selection_reason': details['reason'],
                          'uncertainty_low': lower, 'uncertainty_high': upper,
                          'envelope_hit': lower <= values[cut] <= upper if lower is not None else None})
        details = _selection_details(values)
        details['inner_test_years'] = years[3:]
        method = details['selected_method']
        forecast = _predict(values, method) if len(values) >= 3 else None
        accuracy = _accuracy(tests)
        wape = accuracy['wape_percent']
        baseline_accuracy = _accuracy(fixed_tests['last_year'])
        comparison = [{'method': fixed, **_accuracy(fixed_tests[fixed]),
                       'next_forecast': _predict(values, fixed) if len(values) >= 3 else None,
                       'method_note': _METHOD_NOTES[fixed], 'comparison_note': common['baseline_comparison_note']}
                      for fixed in _METHODS]
        envelope = max((abs(test['error']) for test in tests), default=None)
        judged = [test for test in tests if test['envelope_hit'] is not None]
        hits = sum(test['envelope_hit'] for test in judged)
        reasons = ['Exploratory annual screening only; no operational calibration or vessel-level inference.',
                   'Three-inner-holdout, 5% MAE-margin and two-recent-win promotion controls are heuristic prototype choices.',
                   details['reason']]
        if len(values) < 8:
            reasons.append('Fewer than eight observed annual periods.')
        if len(tests) < 5:
            reasons.append('Fewer than five untouched one-year holdouts.')
        if wape is not None and wape > 15:
            reasons.append('Historical nested-holdout WAPE exceeds the prototype 15% screen; this is not a published accuracy standard.')
        if any(left > 0 and abs(right / left - 1) > .25 for left, right in zip(values[-4:], values[-3:])):
            reasons.append('A recent annual change exceeded the prototype 25% screen.')
        if not tests:
            reasons.append('No untouched outer holdouts available: point prediction exists only when at least three years are supplied, but measured forecast accuracy is unavailable.')
        if tests and baseline_accuracy['mae'] == 0:
            reasons.append('Zero-error last-year benchmark in this small replay; relative improvement is undefined, not a guarantee of future accuracy.')
        if accuracy['mae'] is not None and baseline_accuracy['mae'] is not None:
            if accuracy['mae'] > baseline_accuracy['mae']:
                reasons.append('The guarded policy underperformed the last-year benchmark on these heldout years.')
            elif accuracy['mae'] == baseline_accuracy['mae']:
                reasons.append('The guarded policy matched the last-year benchmark on these heldout years.')
            else:
                reasons.append('The guarded policy improved on the last-year benchmark in this retrospective replay; future improvement is unverified.')
        if tests and wape is None:
            reasons.append('All heldout actual values are zero: WAPE is undefined. MAE and RMSE remain available.')
        output.append({**common, 'forecast': forecast, 'method': method,
                       'status': 'exploratory_not_validated' if forecast is not None else 'insufficient_data',
                       **accuracy, 'tested_first_year': tests[0]['year'] if tests else None,
                       'tested_last_year': tests[-1]['year'] if tests else None,
                       'baseline_wape_percent': baseline_accuracy['wape_percent'], 'baseline_mae': baseline_accuracy['mae'],
                       'baseline_skill_percent': 100 * (1 - accuracy['mae'] / baseline_accuracy['mae']) if baseline_accuracy['mae'] else None,
                       'selection_details': details, 'method_comparison': comparison,
                       'envelope_evaluated_holdouts': len(judged), 'envelope_hits': hits,
                       'envelope_hit_rate_percent': 100 * hits / len(judged) if judged else None,
                       'uncertainty_low': max(0, forecast - envelope) if forecast is not None and envelope is not None else None,
                       'uncertainty_high': forecast + envelope if forecast is not None and envelope is not None else None,
                       'uncertainty_note': 'Largest absolute guarded-policy outer-holdout error around the annual estimate; a historical error envelope, not a confidence interval. It measures past errors of the selection policy, not calibrated coverage for the latest method. Future error can exceed it. Earlier-holdout-only envelope hits are reported separately.' if envelope is not None else 'No outer holdout errors available; uncertainty envelope unavailable.',
                       'reasons': reasons, 'backtests': tests})
    return output


def _json_value(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, allow_nan=False)
    return value


def _csv(path: Path, rows: list[dict], fields: list[str] | None = None):
    columns = fields or list(dict.fromkeys(key for row in rows for key in row))
    with path.open('w', encoding='utf-8-sig', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        for row in rows:
            # Neutralize external text interpreted as spreadsheet formula commands.
            safe = {}
            for key in columns:
                value = _json_value(row.get(key))
                if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
                    value = "'" + value
                safe[key] = value
            writer.writerow(safe)


def _html_table(rows: list[dict], columns: list[str] | None = None) -> str:
    if not rows:
        return '<p>No records available.</p>'
    columns = columns or list(rows[0])
    def show(value):
        if isinstance(value, float):
            return f'{value:,.3f}'
        if value is None:
            return 'unavailable'
        return str(_json_value(value))
    return '<table><thead><tr>' + ''.join(f'<th>{html.escape(col.replace("_", " "))}</th>' for col in columns) + '</tr></thead><tbody>' + ''.join('<tr>' + ''.join(f'<td>{html.escape(show(row.get(col)))}</td>' for col in columns) + '</tr>' for row in rows) + '</tbody></table>'


def dataset_dict(dataset: DatasetBundle) -> dict:
    return {'calls': [asdict(call) for call in dataset.calls],
            'annual_series': [{'port': port, 'metric': metric, 'year': year, 'value': value}
                              for (port, metric), values in sorted(dataset.annual_series.items()) for year, value in values],
            'sources': dataset.sources, 'evidence_notes': dataset.evidence_notes}


def export_report(folder: Path, result: SimulationResult, forecasts: list[dict],
                  recommendations: dict | None, dataset: DatasetBundle) -> list[Path]:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    exported = []
    full = {'schema_version': 1, 'application': 'PortLab Desktop',
            'report_kind': 'conditional_simulation_and_exploratory_annual_forecast',
            'simulation': result.to_dict(), 'forecasts': forecasts,
            'recommendations': recommendations, 'dataset': dataset_dict(dataset)}
    path = folder / 'complete_report.json'
    path.write_text(json.dumps(full, indent=2, ensure_ascii=False, allow_nan=False), encoding='utf-8')
    exported.append(path)
    tables = {
        'vessel_results.csv': (result.vessels, ['vessel_id', 'arrival_hour', 'berth_start_hour', 'departure_hour', 'wait_hours', 'import_boxes', 'export_boxes', 'status']),
        'resource_intervals.csv': (result.intervals, ['resource', 'resource_id', 'start_time', 'end_time', 'vessel_id', 'censored']),
        'stock_flow.csv': (result.kpis.get('stock_flow_samples', []), ['time', 'queued_calls', 'working_calls', 'yard_stock_boxes', 'throughput_boxes', 'pressure_raw', 'pressure_eligible']),
        'event_trace.csv': (result.events, ['time', 'type', 'vessel_id', 'berth', 'box_id', 'flow', 'start_time', 'end_time']),
        'settings.csv': ([{'setting': key, 'value': value} for key, value in result.config.to_dict().items()], None),
        'kpis.csv': ([{'metric': key, 'value': value} for key, value in result.kpis.items() if key != 'stock_flow_samples'], None),
        'checks.csv': ([{'check': key, 'value': value} for key, value in result.checks.items()], None),
        'forecast_results.csv': ([{key: value for key, value in forecast.items() if key not in ('history', 'backtests')} for forecast in forecasts], None),
        'forecast_backtests.csv': ([{'port': forecast['port'], 'metric': forecast['metric'], **test} for forecast in forecasts for test in forecast.get('backtests', [])], None),
        'forecast_methods.csv': ([{'port': forecast['port'], 'metric': forecast['metric'], **row} for forecast in forecasts for row in forecast.get('method_comparison', [])], None),
        'forecast_selection.csv': ([{'port': forecast['port'], 'metric': forecast['metric'], **forecast.get('selection_details', {})} for forecast in forecasts], None),
        'annual_history.csv': (dataset_dict(dataset)['annual_series'], None),
        'input_vessel_calls.csv': ([asdict(call) for call in dataset.calls], None),
        'recommendations.csv': (recommendations.get('candidates', []) if recommendations else [], ['name', 'changes', 'mean_wait_reduction_hours', 'lower_bound_hours', 'served_delta', 'backlog_delta', 'accepted', 'reason', 'paired_replications']),
    }
    for name, (rows, default_fields) in tables.items():
        path = folder / name
        # Preserve every engine field; defaults only define empty-table headers.
        _csv(path, rows, None if rows else default_fields)
        exported.append(path)
    notes = list(dict.fromkeys(dataset.evidence_notes + result.evidence_notes))
    if recommendations:
        notes.append(f'Recommendation status: {recommendations.get("status", "unknown")}. {recommendations.get("reason", "")}')
    content = '''<!doctype html><html><head><meta charset="utf-8"><title>PortLab simulation report</title><style>
    body{font:15px system-ui,sans-serif;color:#172536;max-width:1200px;margin:36px auto;padding:0 22px;line-height:1.5}
    h1,h2{color:#143e58}table{border-collapse:collapse;width:100%;margin:16px 0;font-size:13px}td,th{padding:8px;border:1px solid #d9e1e6;text-align:left;vertical-align:top}th{background:#edf4f7}tr:nth-child(even){background:#f8fafb}.banner{padding:16px;background:#fff2d5;border-left:4px solid #a8751d}.scroll{overflow:auto}code{background:#edf4f7;padding:2px 5px}</style></head><body>'''
    content += f'<h1>{html.escape(result.config.port)} simulation report</h1>'
    content += '<p class="banner">Conditional scenario results. Annual forecasts are exploratory. Inputs, source notes and settings are included below. Incomplete calls remain censored at the horizon. No investment ROI is calculated.</p>'
    content += '<h2>Simulation KPIs</h2>' + _html_table([{'metric': key, 'value': value} for key, value in result.kpis.items() if key != 'stock_flow_samples'])
    content += '<h2>Model checks</h2>' + _html_table([{'check': key, 'value': value} for key, value in result.checks.items()])
    content += '<h2>Annual forecast</h2><p>Point estimates are for the fiscal year following the latest supplied data. Outer errors evaluate the guarded selection policy; the latest selected production method can differ from methods used in earlier folds. This retrospective development replay is not independent prospective validation.</p><div class="scroll">' + _html_table(forecasts, ['port', 'metric', 'history_end_year', 'years', 'forecast_year', 'forecast', 'method', 'wape_percent', 'baseline_wape_percent', 'mae', 'rmse', 'holdouts', 'status', 'uncertainty_low', 'uncertainty_high', 'uncertainty_note']) + '</div>'
    content += '<h2>Actual versus heldout predictions</h2><div class="scroll">' + _html_table(tables['forecast_backtests.csv'][0], ['port', 'metric', 'year', 'train_last_year', 'actual', 'forecast', 'naive_forecast', 'method', 'error', 'absolute_percentage_error', 'uncertainty_low', 'uncertainty_high', 'envelope_hit', 'selection_reason']) + '</div>'
    content += '<h2>Forecast benchmarks</h2><p>Each fixed method is scored on the same outer fiscal years. These descriptive rankings do not select the latest production method. Three inner holdouts, a 5% MAE improvement margin and wins in two recent inner folds are heuristic promotion controls.</p><div class="scroll">' + _html_table(tables['forecast_methods.csv'][0], ['port', 'metric', 'method', 'holdouts', 'mae', 'rmse', 'wape_percent', 'next_forecast', 'method_note']) + '</div>'
    diagnostics = result.kpis.get('operational_prediction_diagnostics')
    if diagnostics:
        content += '<h2>Operational prediction checks</h2>' + _html_table([diagnostics])
    content += '<h2>Scenario comparison</h2><p>' + html.escape(str(recommendations.get('reason') or recommendations.get('status')) if recommendations else 'Scenario comparison not run.') + '</p>'
    content += '<div class="scroll">' + _html_table(recommendations.get('candidates', []) if recommendations else [], ['name', 'changes', 'mean_wait_reduction_hours', 'lower_bound_hours', 'served_delta', 'backlog_delta', 'accepted', 'reason']) + '</div>'
    content += '<h2>Inputs and evidence</h2><ul>' + ''.join('<li>' + html.escape(note) + '</li>' for note in notes) + '</ul>'
    content += '<h2>Settings</h2>' + _html_table([{'setting': key, 'value': value} for key, value in result.config.to_dict().items()])
    content += '<h2>Vessel results</h2><div class="scroll">' + _html_table(result.vessels) + '</div>'
    content += '<h2>Files</h2><ul>' + ''.join(f'<li><a href="{html.escape(path.name)}">{html.escape(path.name)}</a></li>' for path in exported) + '</ul></body></html>'
    path = folder / 'report.html'
    path.write_text(content, encoding='utf-8')
    exported.append(path)
    # Runtime exporter in the delivered Python application, not authored analysis.
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        book = openpyxl.Workbook()
        book.remove(book.active)
        for name, (rows, default_fields) in tables.items():
            if name == 'event_trace.csv':
                continue  # Full trace is in CSV/JSON; keep the workbook readable.
            sheet = book.create_sheet(name.removesuffix('.csv')[:31])
            fields = list(dict.fromkeys(key for row in rows for key in row)) or default_fields or ['status']
            sheet.append(fields)
            for row in rows:
                sheet.append([_json_value(row.get(key)) for key in fields])
                for cell in sheet[sheet.max_row]:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='143E58')
            for index, field in enumerate(fields, 1):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = min(55, max(14, len(field) + 3))
        sheet = book.create_sheet('Evidence notes')
        sheet.append(['Evidence note'])
        for note in notes:
            sheet.append([note])
            sheet.cell(sheet.max_row, 1).data_type = 's'
            sheet.cell(sheet.max_row, 1).alignment = Alignment(wrap_text=True, vertical='top')
            sheet.row_dimensions[sheet.max_row].height = max(30, math.ceil(len(note) / 95) * 16)
        sheet.column_dimensions['A'].width = 110
        path = folder / 'results.xlsx'
        book.save(path)
        exported.append(path)
    except ImportError:
        pass
    if (folder / 'results.xlsx') in exported:
        report_path = folder / 'report.html'
        report_text = report_path.read_text(encoding='utf-8')
        report_path.write_text(report_text.replace('</ul></body></html>',
                                                 '<li><a href="results.xlsx">results.xlsx</a></li></ul></body></html>'),
                               encoding='utf-8')
    return exported
